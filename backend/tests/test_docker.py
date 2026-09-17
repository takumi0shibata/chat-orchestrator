import asyncio
import os
from uuid import uuid4

import pytest
from app.config import Folder, Settings, Skill
from app.sandbox import Sandbox, docker

pytestmark = [
    pytest.mark.docker,
    pytest.mark.skipif(
        os.environ.get("RUN_DOCKER_TESTS") != "1", reason="Set RUN_DOCKER_TESTS=1"
    ),
]


async def emit(*args):
    pass


def test_200_files_office_pdf_research_and_isolation(tmp_path):
    async def scenario():
        work = tmp_path / "workspace"
        work.mkdir()
        inputs = tmp_path / "input"
        inputs.mkdir()
        skill_path = tmp_path / "skill"
        skill_path.mkdir()
        (skill_path / "SKILL.md").write_text(
            "---\nname: test\ndescription: Test\n---\nUse local files."
        )
        (tmp_path / "host-secret").write_text("DO NOT READ")
        (work / "escape").symlink_to(tmp_path / "host-secret")
        for i in range(195):
            (work / f"file-{i}.txt").write_text(f"original {i}")
        sandbox = Sandbox(
            Settings(_env_file=None),
            Folder(id="w", label="w", path=work),
            uuid4().hex,
            [Skill(id="test", label="test", path=skill_path)],
            [],
            inputs,
        )
        await sandbox.start()
        try:
            script = r"""
import os, socket
from pathlib import Path
import numpy as np, pandas as pd, polars as pl, scipy, torch, sklearn
from sklearn.linear_model import LinearRegression
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import seaborn as sns
import transformers, datasets
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader
assert os.getuid() != 0
assert not torch.cuda.is_available()
assert torch.version.cuda is None
assert not Path('/var/run/docker.sock').exists()
assert not Path('/workspace/escape').exists()
assert not os.environ.get('OPENAI_API_KEY')
assert not os.environ.get('AZURE_OPENAI_API_KEY')
try:
    Path('/skills/test/SKILL.md').write_text('overwrite')
    raise AssertionError('skill was writable')
except OSError: pass
try:
    Path('/input/test.txt').write_text('overwrite')
    raise AssertionError('input was writable')
except OSError: pass
try:
    socket.create_connection(('1.1.1.1',443),timeout=1)
    raise AssertionError('network was available')
except OSError: pass
Document().save('paper.docx')
doc=Document('paper.docx'); doc.add_paragraph('edited original'); doc.save('paper.docx')
wb=Workbook(); wb.active['A1']=42; wb.save('table.xlsx'); assert load_workbook('table.xlsx').active['A1'].value==42
prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save('slides.pptx'); assert len(Presentation('slides.pptx').slides)==1
c=canvas.Canvas('paper.pdf'); c.drawString(40,750,'PDF test'); c.save(); assert 'PDF test' in PdfReader('paper.pdf').pages[0].extract_text()
pd.DataFrame({'x':[1,2,3],'y':[2,4,6]}).to_csv('data.csv',index=False)
assert len([p for p in Path('.').iterdir() if p.is_file()])==200
Path('file-194.txt').write_text('direct edit')
assert torch.tensor([1.,2.]).sum().item()==3
assert np.isclose(LinearRegression().fit([[1],[2],[3]],[2,4,6]).predict([[4]])[0],8)
plt.rcParams['font.family']='Noto Sans CJK JP'
assert font_manager.findfont('Noto Sans CJK JP',fallback_to_default=False)
sns.lineplot(x=[1,2,3],y=[2,4,6]); plt.title('日本語の分析結果'); plt.savefig('analysis.png'); plt.close()
print('200 mixed files + direct edit + Office/PDF + CPU ML + Japanese plot + isolation OK')
"""
            result = await sandbox.execute(
                "python - <<'PY'\n" + script + "\nPY", emit, 180
            )
            assert result["outcome"]["exit_code"] == 0, result
            assert "isolation OK" in result["stdout"]
            assert (work / "file-194.txt").read_text() == "direct edit"
            assert (work / "analysis.png").stat().st_size > 1000
            # Actual LibreOffice conversion works without network or a writable root filesystem.
            converted = await sandbox.execute(
                "mkdir -p converted && libreoffice -env:UserInstallation=file:///tmp/lo-test --headless --convert-to pdf --outdir converted paper.docx",
                emit,
                60,
            )
            assert converted["outcome"]["exit_code"] == 0, converted
            assert (work / "converted/paper.pdf").exists()
        finally:
            await sandbox.close()
        assert (work / "file-194.txt").read_text() == "direct edit"

    asyncio.run(scenario())


def test_timeout_removes_container_and_descendants(tmp_path):
    async def scenario():
        s = Settings(_env_file=None, command_timeout=1)
        sandbox = Sandbox(
            s,
            Folder(id="w", label="w", path=tmp_path),
            uuid4().hex,
            [],
            [],
            tmp_path / "input",
        )
        await sandbox.start()
        with pytest.raises(TimeoutError):
            await sandbox.execute("sleep 100 & wait", emit)
        with pytest.raises(RuntimeError):
            await docker("inspect", sandbox.name)
        await sandbox.close()

    asyncio.run(scenario())


def test_cancel_removes_container(tmp_path):
    async def scenario():
        sandbox = Sandbox(
            Settings(_env_file=None),
            Folder(id="w", label="w", path=tmp_path),
            uuid4().hex,
            [],
            [],
            tmp_path / "input",
        )
        await sandbox.start()
        task = asyncio.create_task(sandbox.execute("sleep 100 & wait", emit))
        await asyncio.sleep(0.5)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        with pytest.raises(RuntimeError):
            await docker("inspect", sandbox.name)

    asyncio.run(scenario())
